"""
Este arquivo cuida de IMPORTAR dados de uma planilha .xlsx (Excel) pra dentro
de um banco de dados novo -- e usado logo depois que o usuario cria um banco
pela primeira vez, oferecendo trazer os dados que ja existiam (por exemplo,
do arquivo "Presidentes - Mailing.xlsx").

Regra importante: o programa le o CABECALHO de cada aba da planilha (a
primeira linha, com os nomes das colunas) na hora, em vez de ja "saber de
cor" quais colunas cada aba tem. Isso significa que, se a planilha mudar um
pouco no futuro (colunas a mais, a menos, em outra ordem), a importacao
continua funcionando sem precisar mexer no codigo do programa.

Como cada aba da planilha e tratada
------------------------------------
- A aba chamada "EMPRESAS" entra direto na tabela EMPRESAS do banco novo,
  mantendo o mesmo numero de ID que tinha na planilha -- isso e importante
  porque as abas de pessoas apontam pra esses IDs.
- Qualquer OUTRA aba que tenha uma coluna chamada "ID_EMPRESA" e considerada
  uma "aba de gente" (presidentes, diretores etc.) e todas elas sao
  misturadas numa unica tabela PESSOAS -- e essa a mudanca de modelo que o
  usuario pediu: em vez de uma tabela separada por cargo, uma tabela so, com
  um campo CARGO dizendo o papel de cada pessoa.
- Abas sem ID_EMPRESA (que nao sejam EMPRESAS) viram tabelas novas e
  independentes, exatamente como vieram na planilha.

Alem disso, dois ajustes de qualidade sao feitos automaticamente durante a
importacao (sem alterar a planilha original, so os dados que entram no
banco novo):
- CATEGORIA: cada pessoa ganha uma etiqueta com o nome da aba de onde ela
  veio (ex.: "Presidentes", "Diretores Tecnicos") -- diferente do campo
  CARGO (texto livre, digitado na planilha), essa etiqueta e o que permite
  reconstruir a mala direta "uma linha por empresa" que a planilha antiga
  fazia com as abas separadas (ver db/exporter.py).
- Nomes proprios (de pessoa e de empresa) que vem em CAIXA ALTA sao
  normalizados pra so a primeira letra de cada palavra maiuscula (ver
  db/normalizacao.py) -- a coluna EMPRESA aproveita a coluna
  EMPRESA_MINUSCULA da planilha original quando ela existe, que ja vem
  com a capitalizacao certa.
"""
from __future__ import annotations

import datetime as _dt
import sqlite3
from dataclasses import dataclass, field

import openpyxl

from . import log, tables
from .identifiers import quote_ident, validar_identificador
from .normalizacao import normalizar_nome_proprio
from .schema import EMPRESAS, PESSOAS

# Abas que nunca devem ser importadas -- ou porque nao fazem mais sentido
# nesse formato (CONFIGURACOES virou a tabela app_log) ou porque o banco
# novo ja nasce com a tabela USUARIOS propria (a senha da planilha antiga,
# se existisse, nao seria compativel com o novo sistema de hash+pepper).
_ABAS_IGNORADAS = {"CONFIGURAÇÕES", "CONFIGURACOES", "USUARIOS"}

# Campos de nome proprio que vem em CAIXA ALTA na planilha original e sao
# normalizados (so a primeira letra de cada palavra maiuscula) na hora de
# importar -- ver db/normalizacao.py.
_CAMPOS_NOME_PROPRIO_EMPRESA = ["EMPRESA"]
_CAMPOS_NOME_PROPRIO_PESSOA = ["NOME", "NOME DE GUERRA", "ASSESSOR(A)"]


@dataclass
class ResumoImportacao:
    """Guarda o resultado da importacao, pra mostrar um resumo pro usuario
    no final (quantas linhas entraram em cada aba, e avisos se alguma aba
    veio vazia)."""
    linhas_por_aba: dict = field(default_factory=dict)
    avisos: list = field(default_factory=list)


def _valor_celula(v):
    """Converte o valor "cru" que o Excel guarda numa celula pro formato que
    o banco de dados espera. Datas viram texto no formato AAAA-MM-DD (um
    padrao internacional que ordena certinho); textos tem espacos nas pontas
    removidos; o resto (numeros etc.) passa direto."""
    if isinstance(v, (_dt.datetime, _dt.date)):
        return v.date().isoformat() if isinstance(v, _dt.datetime) else v.isoformat()
    if isinstance(v, str):
        v = v.strip()
        return v if v else None
    return v


def _ler_cabecalho(ws) -> list[str]:
    """Le a primeira linha da aba (os titulos das colunas). Se um titulo
    aparecer repetido (ex.: duas colunas chamadas "TELEFONE"), a segunda
    ocorrencia ganha um sufixo " (2)" pra nao se confundir com a primeira."""
    primeira_linha = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
    vistos: dict[str, int] = {}
    cabecalho = []
    for valor in primeira_linha:
        nome = str(valor).strip() if valor is not None else ""
        if not nome:
            cabecalho.append("")  # coluna sem titulo -- sera ignorada
            continue
        vistos[nome] = vistos.get(nome, 0) + 1
        cabecalho.append(nome if vistos[nome] == 1 else f"{nome} ({vistos[nome]})")
    return cabecalho


def _ler_linhas(ws, cabecalho: list[str]) -> list[dict]:
    """Le todas as linhas de dados (a partir da segunda linha, ja que a
    primeira e o cabecalho), devolvendo cada uma como um dicionario
    {nome_da_coluna: valor}. Linhas completamente vazias sao descartadas."""
    linhas = []
    for linha_bruta in ws.iter_rows(min_row=2, values_only=True):
        registro = {}
        for coluna, valor in zip(cabecalho, linha_bruta):
            if not coluna:
                continue  # coluna sem titulo no cabecalho -- ignora essa celula
            registro[coluna] = _valor_celula(valor)
        if any(v not in (None, "") for v in registro.values()):
            linhas.append(registro)
    return linhas


def _inserir_linha(conn: sqlite3.Connection, tabela: str, dados: dict) -> None:
    """Monta e executa um INSERT simples com os campos de `dados`."""
    campos = list(dados.keys())
    if not campos:
        conn.execute(f"INSERT INTO {quote_ident(tabela)} DEFAULT VALUES")
        return
    campos_sql = ", ".join(quote_ident(c) for c in campos)
    placeholders = ", ".join("?" for _ in campos)
    conn.execute(
        f"INSERT INTO {quote_ident(tabela)} ({campos_sql}) VALUES ({placeholders})",
        [dados[c] for c in campos],
    )


def _garantir_colunas(conn: sqlite3.Connection, tabela: str, colunas_desejadas: list[str], usuario: str) -> None:
    """Antes de importar, confere se a tabela de destino ja tem todas as
    colunas que a planilha traz -- se faltar alguma, cria na hora (usando o
    esquema dinamico de db/tables.py). Assim a importacao nunca perde dados
    por causa de uma coluna que o banco "nao esperava"."""
    existentes = set(tables.get_schema(conn, tabela))
    for coluna in colunas_desejadas:
        if coluna and coluna != "ID" and coluna not in existentes:
            tables.add_column(conn, tabela, coluna, usuario=usuario)
            existentes.add(coluna)


def _importar_empresas(conn: sqlite3.Connection, cabecalho: list[str], linhas: list[dict], usuario: str) -> int:
    """Importa a aba EMPRESAS, mantendo o ID original de cada empresa (as
    pessoas importadas depois vao apontar pra esses mesmos IDs)."""
    _garantir_colunas(conn, EMPRESAS, cabecalho, usuario)
    for linha in linhas:
        # A planilha original ja tinha uma coluna "EMPRESA_MINUSCULA" com o
        # nome da empresa formatado corretamente (so a primeira letra de
        # cada palavra maiuscula) -- reaproveitamos ela pra corrigir a
        # coluna EMPRESA (que vem em CAIXA ALTA), em vez de tentar adivinhar
        # a capitalizacao certa do zero.
        if linha.get("EMPRESA_MINUSCULA"):
            linha["EMPRESA"] = linha["EMPRESA_MINUSCULA"]
        else:
            for campo in _CAMPOS_NOME_PROPRIO_EMPRESA:
                if linha.get(campo):
                    linha[campo] = normalizar_nome_proprio(linha[campo])
        _inserir_linha(conn, EMPRESAS, linha)
    conn.commit()
    return len(linhas)


def _importar_pessoas(conn: sqlite3.Connection, aba_origem: str, cabecalho: list[str],
                       linhas: list[dict], usuario: str) -> int:
    """Importa uma aba de pessoas (qualquer aba com coluna ID_EMPRESA) pra
    dentro da tabela unica PESSOAS.

    O ID original de cada linha na planilha e DESCARTADO de proposito: como
    varias abas antigas (Presidentes, Diretores...) tinham cada uma sua
    propria contagem de ID comecando do 1, misturar todas numa tabela so exige
    gerar numeros novos, senao teriamos IDs repetidos. O ID_EMPRESA (que
    aponta pra outra tabela) esse sim e mantido, porque continua valendo.
    """
    colunas_pessoa = [c for c in cabecalho if c != "ID"]
    if "CATEGORIA" not in colunas_pessoa:
        colunas_pessoa.append("CATEGORIA")
    _garantir_colunas(conn, PESSOAS, colunas_pessoa, usuario)
    categoria = normalizar_nome_proprio(aba_origem)

    for linha in linhas:
        dados = {k: v for k, v in linha.items() if k != "ID"}

        # CATEGORIA guarda de qual aba a pessoa veio (ex.: "Presidentes",
        # "Diretores Tecnicos") -- diferente de CARGO, que e o titulo livre
        # digitado na planilha (ex.: "Secretario de Estado..."). Sem isso,
        # a informacao de "aba de origem" da planilha antiga se perderia.
        dados["CATEGORIA"] = categoria
        # Se a linha nao tinha um CARGO preenchido, usamos a categoria como
        # cargo tambem -- assim nenhuma pessoa fica sem indicacao de papel.
        if not dados.get("CARGO"):
            dados["CARGO"] = categoria

        for campo in _CAMPOS_NOME_PROPRIO_PESSOA:
            if dados.get(campo):
                dados[campo] = normalizar_nome_proprio(dados[campo])

        _inserir_linha(conn, PESSOAS, dados)
    conn.commit()
    return len(linhas)


def _importar_tabela_generica(conn: sqlite3.Connection, nome_aba: str, cabecalho: list[str],
                               linhas: list[dict], usuario: str) -> int:
    """Importa uma aba que nao e EMPRESAS nem tem ID_EMPRESA, criando uma
    tabela nova e independente com o mesmo nome e colunas da aba."""
    nome = validar_identificador(nome_aba, "nome da tabela")
    existentes = {t.lower() for t in tables.list_data_sheets(conn)}
    if nome.lower() not in existentes:
        tables.create_data_sheet(conn, nome, usuario=usuario)
    _garantir_colunas(conn, nome, cabecalho, usuario)
    for linha in linhas:
        _inserir_linha(conn, nome, linha)
    conn.commit()
    return len(linhas)


def importar_xlsx(conn: sqlite3.Connection, caminho_xlsx: str, usuario: str = "sistema") -> ResumoImportacao:
    """Funcao principal: le o arquivo .xlsx informado e importa todas as
    abas relevantes pro banco de dados `conn` (que ja deve existir, com o
    schema inicial criado)."""
    resumo = ResumoImportacao()
    wb = openpyxl.load_workbook(caminho_xlsx, data_only=True, read_only=True)
    try:
        planilhas = list(wb.worksheets)
        # A aba EMPRESAS precisa ser importada ANTES de qualquer aba de
        # pessoas, senao o ID_EMPRESA das pessoas apontaria pra uma empresa
        # que ainda nao existe no banco novo.
        planilhas.sort(key=lambda ws: 0 if ws.title.strip().upper() == EMPRESAS else 1)

        for ws in planilhas:
            nome_aba = ws.title.strip()
            if nome_aba.upper() in _ABAS_IGNORADAS:
                continue

            cabecalho = _ler_cabecalho(ws)
            if not any(cabecalho):
                continue  # aba sem nenhum titulo de coluna -- nada a importar
            linhas = _ler_linhas(ws, cabecalho)
            if not linhas:
                resumo.avisos.append(f'Aba "{nome_aba}" nao tinha linhas de dados.')
                continue

            if nome_aba.upper() == EMPRESAS:
                qtd = _importar_empresas(conn, cabecalho, linhas, usuario)
            elif "ID_EMPRESA" in cabecalho:
                qtd = _importar_pessoas(conn, nome_aba, cabecalho, linhas, usuario)
            else:
                qtd = _importar_tabela_generica(conn, nome_aba, cabecalho, linhas, usuario)

            resumo.linhas_por_aba[nome_aba] = qtd
            log.log_change(conn, usuario, nome_aba, "Importacao inicial", f"{qtd} linha(s)")
    finally:
        wb.close()

    return resumo
