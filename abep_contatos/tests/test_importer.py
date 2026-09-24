import os

import pytest

from db import importer, records, tables

_XLSX_REAL = os.path.join(
    os.path.dirname(__file__), "..", "..", "Presidentes - Mailing.xlsx"
)


@pytest.mark.skipif(not os.path.isfile(_XLSX_REAL), reason="planilha real nao encontrada")
def test_importar_planilha_real_funde_pessoas(conn):
    resumo = importer.importar_xlsx(conn, _XLSX_REAL)

    empresas = records.get_records(conn, "EMPRESAS")
    assert len(empresas) > 0

    pessoas = records.get_records(conn, "PESSOAS")
    # soma das linhas de todas as abas de pessoas fundidas
    total_pessoas_no_resumo = sum(
        qtd for aba, qtd in resumo.linhas_por_aba.items() if aba.upper() != "EMPRESAS"
    )
    assert len(pessoas) == total_pessoas_no_resumo
    assert len(pessoas) > 0

    # toda pessoa deve ter ID_EMPRESA resolvendo pra uma empresa real
    ids_empresas = {e["ID"] for e in empresas}
    for pessoa in pessoas:
        assert pessoa["ID_EMPRESA"] in ids_empresas
        assert pessoa.get("_EMPRESA_SIGLA")

    # cada pessoa tem um CARGO preenchido (herdado da planilha ou da aba de origem)
    assert all(p.get("CARGO") for p in pessoas)


def test_importar_xlsx_sintetico(tmp_path, conn):
    import openpyxl

    caminho = tmp_path / "teste.xlsx"
    wb = openpyxl.Workbook()

    ws_empresas = wb.active
    ws_empresas.title = "EMPRESAS"
    ws_empresas.append(["ID", "SIGLA", "EMPRESA"])
    ws_empresas.append([1, "ABC", "Empresa ABC"])
    ws_empresas.append([2, "XYZ", "Empresa XYZ"])

    ws_pres = wb.create_sheet("PRESIDENTES")
    ws_pres.append(["ID", "ID_EMPRESA", "CARGO", "NOME"])
    ws_pres.append([1, 1, "Presidente", "Ana"])
    ws_pres.append([2, 2, "Presidente", "Bruno"])

    ws_dir = wb.create_sheet("DIRETORES TÉCNICOS")
    ws_dir.append(["ID", "ID_EMPRESA", "CARGO", "NOME"])
    ws_dir.append([1, 1, "Diretor Técnico", "Carla"])

    wb.save(caminho)

    resumo = importer.importar_xlsx(conn, str(caminho))

    assert resumo.linhas_por_aba["EMPRESAS"] == 2
    assert resumo.linhas_por_aba["PRESIDENTES"] == 2
    assert resumo.linhas_por_aba["DIRETORES TÉCNICOS"] == 1

    pessoas = records.get_records(conn, "PESSOAS")
    assert len(pessoas) == 3
    nomes = {p["NOME"] for p in pessoas}
    assert nomes == {"Ana", "Bruno", "Carla"}

    # IDs de PESSOAS sao gerados de novo (sequenciais), nao reaproveitados das abas
    assert sorted(p["ID"] for p in pessoas) == [1, 2, 3]


def test_importar_tabela_generica_sem_id_empresa(tmp_path, conn):
    import openpyxl

    caminho = tmp_path / "teste2.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "FORNECEDORES"
    ws.append(["ID", "NOME", "TELEFONE"])
    ws.append([1, "Fornecedor X", "1111-1111"])
    wb.save(caminho)

    importer.importar_xlsx(conn, str(caminho))

    assert "FORNECEDORES" in tables.list_data_sheets(conn)
    fornecedores = records.get_records(conn, "FORNECEDORES")
    assert len(fornecedores) == 1
    assert fornecedores[0]["NOME"] == "Fornecedor X"
