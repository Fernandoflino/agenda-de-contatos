from db import categorias, exporter, records


def _preparar(conn):
    id_abc = records.create_record(conn, "EMPRESAS", {"SIGLA": "ABC", "EMPRESA": "Empresa ABC"})
    id_xyz = records.create_record(conn, "EMPRESAS", {"SIGLA": "XYZ", "EMPRESA": "Empresa XYZ"})

    records.create_record(conn, "PESSOAS", {"ID_EMPRESA": id_abc, "CARGO": "Presidente", "NOME": "Ana"})
    records.create_record(conn, "PESSOAS", {"ID_EMPRESA": id_abc, "CARGO": "Diretor Técnico", "NOME": "Carla"})
    records.create_record(conn, "PESSOAS", {"ID_EMPRESA": id_xyz, "CARGO": "Presidente", "NOME": "Bruno"})
    # empresa XYZ tem 2 diretores tecnicos -- caso que motivou abandonar as 3 tabelas fixas
    records.create_record(conn, "PESSOAS", {"ID_EMPRESA": id_xyz, "CARGO": "Diretor Técnico", "NOME": "Duda"})
    records.create_record(conn, "PESSOAS", {"ID_EMPRESA": id_xyz, "CARGO": "Diretor Técnico", "NOME": "Elis"})
    return id_abc, id_xyz


def test_exportacao_simples_filtra_por_cargo(conn):
    _preparar(conn)
    columns, rows = exporter.montar_exportacao_simples(
        conn, "PESSOAS", filtro_campo="CARGO", filtro_valor="Diretor Técnico"
    )
    assert len(rows) == 3
    assert all(r["CARGO"] == "Diretor Técnico" for r in rows)
    assert "NOME" in columns


def test_exportacao_simples_restringe_por_ids_permitidos(conn):
    """"Exportar selecionados" (acao em massa na lista de registros) precisa
    exportar SO os IDs marcados, mesmo sem nenhum outro filtro de campo."""
    _preparar(conn)
    todos = records.get_records(conn, "PESSOAS")
    ids_selecionados = {todos[0]["ID"], todos[2]["ID"]}

    columns, rows = exporter.montar_exportacao_simples(conn, "PESSOAS", ids_permitidos=ids_selecionados)
    assert len(rows) == 2
    assert {r["NOME"] for r in rows} == {todos[0]["NOME"], todos[2]["NOME"]}


def test_colunas_exportaveis_nao_inclui_foto(conn):
    """FOTO/FOTO_MIME sao BLOB -- nao fazem sentido numa celula de
    planilha, entao nunca aparecem como coluna exportavel."""
    colunas, _ = exporter.colunas_exportaveis(conn, "PESSOAS")
    assert "FOTO" not in colunas
    assert "FOTO_MIME" not in colunas
    assert "FOTO_ORIGINAL" not in colunas
    assert "FOTO_ORIGINAL_MIME" not in colunas


def test_exportacao_simples_mostra_empresa_em_vez_de_id(conn):
    _preparar(conn)
    columns, rows = exporter.montar_exportacao_simples(conn, "PESSOAS")
    assert "ID_EMPRESA" not in columns
    assert "SIGLA" in columns
    assert "EMPRESA" in columns
    assert {r["SIGLA"] for r in rows} == {"ABC", "XYZ"}
    assert {r["EMPRESA"] for r in rows} == {"Empresa ABC", "Empresa XYZ"}


def test_exportacao_simples_empresas_mantem_suas_proprias_colunas(conn):
    _preparar(conn)
    columns, rows = exporter.montar_exportacao_simples(conn, "EMPRESAS")
    assert "SIGLA" in columns
    assert "EMPRESA" in columns
    siglas = {r["SIGLA"] for r in rows}
    assert siglas == {"ABC", "XYZ"}


def test_exportacao_mesclada_uma_linha_por_empresa_simples(conn):
    _preparar(conn)
    columns, rows = exporter.montar_exportacao_mesclada(
        conn, "PESSOAS", valores_agrupador=["Presidente"], campo_agrupador="CARGO",
        campos_por_valor={"Presidente": ["NOME"]},
    )
    # so 1 presidente por empresa -> 1 linha por empresa, 2 empresas
    assert len(rows) == 2
    nomes_presidentes = {r["Presidente - NOME"] for r in rows}
    assert nomes_presidentes == {"Ana", "Bruno"}


def test_exportacao_mesclada_repete_linha_quando_ha_mais_de_uma_pessoa_no_cargo(conn):
    _preparar(conn)
    columns, rows = exporter.montar_exportacao_mesclada(
        conn, "PESSOAS", campo_agrupador="CARGO",
        valores_agrupador=["Presidente", "Diretor Técnico"],
        campos_por_valor={"Presidente": ["NOME"], "Diretor Técnico": ["NOME"]},
    )
    # ABC: 1 presidente x 1 diretor = 1 linha. XYZ: 1 presidente x 2 diretores = 2 linhas.
    assert len(rows) == 3
    nomes_diretores_xyz = {
        r["Diretor Técnico - NOME"] for r in rows if r["Presidente - NOME"] == "Bruno"
    }
    assert nomes_diretores_xyz == {"Duda", "Elis"}


def test_exportacao_mesclada_usa_categoria_por_padrao(conn):
    """A categoria (etiqueta que guarda de qual aba antiga a pessoa veio) e
    o agrupador padrao -- diferente do CARGO (texto livre), ela tem sempre
    os mesmos valores conhecidos, o que a torna mais confiavel pra montar a
    mala direta."""
    categorias.adicionar_categoria(conn, "Presidentes")
    categorias.adicionar_categoria(conn, "Diretores Tecnicos")
    id_abc = records.create_record(conn, "EMPRESAS", {"SIGLA": "QQQ", "EMPRESA": "Empresa QQQ"})
    id_fabio = records.create_record(conn, "PESSOAS", {
        "ID_EMPRESA": id_abc, "CARGO": "Presidente Executivo", "NOME": "Fabio",
    })
    id_gina = records.create_record(conn, "PESSOAS", {
        "ID_EMPRESA": id_abc, "CARGO": "Coordenador-Geral", "NOME": "Gina",
    })
    categorias.definir_categorias_da_pessoa(conn, id_fabio, ["Presidentes"])
    categorias.definir_categorias_da_pessoa(conn, id_gina, ["Diretores Tecnicos"])

    columns, rows = exporter.montar_exportacao_mesclada(
        conn, "PESSOAS", valores_agrupador=["Presidentes", "Diretores Tecnicos"],
        campos_por_valor={"Presidentes": ["NOME"], "Diretores Tecnicos": ["NOME"]},
    )
    assert len(rows) == 1
    assert rows[0]["Presidentes - NOME"] == "Fabio"
    assert rows[0]["Diretores Tecnicos - NOME"] == "Gina"


def test_exportacao_mesclada_pessoa_com_varias_categorias_aparece_em_todos_os_blocos(conn):
    """Um contato pode ter mais de uma categoria ao mesmo tempo -- ele deve
    aparecer em CADA bloco correspondente, nao so num."""
    categorias.adicionar_categoria(conn, "Presidentes")
    categorias.adicionar_categoria(conn, "Diretores Tecnicos")
    id_empresa = records.create_record(conn, "EMPRESAS", {"SIGLA": "QQQ", "EMPRESA": "Empresa QQQ"})
    id_fabio = records.create_record(conn, "PESSOAS", {"ID_EMPRESA": id_empresa, "NOME": "Fabio"})
    categorias.definir_categorias_da_pessoa(conn, id_fabio, ["Presidentes", "Diretores Tecnicos"])

    columns, rows = exporter.montar_exportacao_mesclada(
        conn, "PESSOAS", valores_agrupador=["Presidentes", "Diretores Tecnicos"],
        campos_por_valor={"Presidentes": ["NOME"], "Diretores Tecnicos": ["NOME"]},
    )
    assert len(rows) == 1
    assert rows[0]["Presidentes - NOME"] == "Fabio"
    assert rows[0]["Diretores Tecnicos - NOME"] == "Fabio"


def test_exportacao_simples_junta_varias_categorias_com_virgula(conn):
    categorias.adicionar_categoria(conn, "Presidentes")
    categorias.adicionar_categoria(conn, "Diretores Tecnicos")
    id_pessoa = records.create_record(conn, "PESSOAS", {"NOME": "Fabio"})
    categorias.definir_categorias_da_pessoa(conn, id_pessoa, ["Presidentes", "Diretores Tecnicos"])

    columns, rows = exporter.montar_exportacao_simples(conn, "PESSOAS")

    assert "CATEGORIA" in columns
    linha = next(r for r in rows if r["NOME"] == "Fabio")
    assert linha["CATEGORIA"] == "Presidentes, Diretores Tecnicos"


def test_exportacao_mesclada_nao_inclui_foto(conn):
    """FOTO/FOTO_MIME/FOTO_ORIGINAL/FOTO_ORIGINAL_MIME sao BLOB -- assim como
    no modo Simples (test_colunas_exportaveis_nao_inclui_foto), o modo
    Mesclado nao pode incluir esses campos crus num bloco de encaixe: alem
    de nao fazerem sentido numa celula, quebravam a exportacao (bytes crus
    de imagem tentados como texto)."""
    id_empresa = records.create_record(conn, "EMPRESAS", {"SIGLA": "ABC", "EMPRESA": "Empresa ABC"})
    records.create_record(conn, "PESSOAS", {
        "ID_EMPRESA": id_empresa, "CARGO": "Presidente", "NOME": "Ana",
        "FOTO": b"\x89PNG\r\n\x1a\n...", "FOTO_MIME": "image/png",
        "FOTO_ORIGINAL": b"\x89PNG\r\n\x1a\n...", "FOTO_ORIGINAL_MIME": "image/png",
    })

    columns, rows = exporter.montar_exportacao_mesclada(
        conn, "PESSOAS", valores_agrupador=["Presidente"], campo_agrupador="CARGO",
    )

    assert not any(c.endswith(" - FOTO") for c in columns)
    assert not any(c.endswith(" - FOTO_MIME") for c in columns)
    assert not any(c.endswith(" - FOTO_ORIGINAL") for c in columns)
    assert not any(c.endswith(" - FOTO_ORIGINAL_MIME") for c in columns)
    assert not any(isinstance(v, (bytes, bytearray)) for r in rows for v in r.values())


def test_exportacao_mesclada_com_foto_nao_quebra_xlsx_e_csv(conn, tmp_path):
    """Regressao: exportar o modo Mesclado com gente que tem foto cadastrada
    lancava UnicodeDecodeError (bytes crus de imagem indo pra uma celula)."""
    id_empresa = records.create_record(conn, "EMPRESAS", {"SIGLA": "ABC", "EMPRESA": "Empresa ABC"})
    records.create_record(conn, "PESSOAS", {
        "ID_EMPRESA": id_empresa, "CARGO": "Presidente", "NOME": "Ana",
        "FOTO": b"\x89PNG\r\n\x1a\n...", "FOTO_MIME": "image/png",
    })

    columns, rows = exporter.montar_exportacao_mesclada(
        conn, "PESSOAS", valores_agrupador=["Presidente"], campo_agrupador="CARGO",
    )

    exporter.exportar_xlsx(columns, rows, str(tmp_path / "mesclado.xlsx"))
    exporter.exportar_csv(columns, rows, str(tmp_path / "mesclado.csv"))
    assert (tmp_path / "mesclado.xlsx").is_file()
    assert (tmp_path / "mesclado.csv").is_file()


def test_exportar_xlsx_e_csv(conn, tmp_path):
    _preparar(conn)
    columns, rows = exporter.montar_exportacao_simples(conn, "PESSOAS")

    caminho_xlsx = tmp_path / "export.xlsx"
    exporter.exportar_xlsx(columns, rows, str(caminho_xlsx))
    assert caminho_xlsx.is_file()

    caminho_csv = tmp_path / "export.csv"
    exporter.exportar_csv(columns, rows, str(caminho_csv))
    conteudo = caminho_csv.read_text(encoding="utf-8-sig")
    assert "Ana" in conteudo


def test_nome_coluna_arquivo_troca_espaco_e_traco_por_underscore():
    assert exporter.nome_coluna_arquivo("DATA DE NASCIMENTO") == "DATA_DE_NASCIMENTO"
    assert exporter.nome_coluna_arquivo("Presidente - NOME") == "Presidente_NOME"
    assert exporter.nome_coluna_arquivo("NOME") == "NOME"


def test_exportar_xlsx_troca_espaco_por_underscore_no_cabecalho(conn, tmp_path):
    id_empresa = records.create_record(conn, "EMPRESAS", {"SIGLA": "ABC", "EMPRESA": "Empresa ABC"})
    records.create_record(conn, "PESSOAS", {
        "ID_EMPRESA": id_empresa, "NOME": "Ana", "DATA DE NASCIMENTO": "2000-01-01",
    })
    columns, rows = exporter.montar_exportacao_simples(conn, "PESSOAS")

    caminho = tmp_path / "export.xlsx"
    exporter.exportar_xlsx(columns, rows, str(caminho))

    import openpyxl
    wb = openpyxl.load_workbook(str(caminho))
    cabecalho = [c.value for c in next(wb.active.iter_rows(min_row=1, max_row=1))]
    assert "DATA_DE_NASCIMENTO" in cabecalho
    assert "DATA DE NASCIMENTO" not in cabecalho

    linha = next(wb.active.iter_rows(min_row=2, max_row=2, values_only=True))
    indice_data = cabecalho.index("DATA_DE_NASCIMENTO")
    assert linha[indice_data] == "2000-01-01"


def test_exportar_csv_troca_espaco_por_underscore_no_cabecalho(conn, tmp_path):
    id_empresa = records.create_record(conn, "EMPRESAS", {"SIGLA": "ABC", "EMPRESA": "Empresa ABC"})
    records.create_record(conn, "PESSOAS", {
        "ID_EMPRESA": id_empresa, "NOME": "Ana", "DATA DE NASCIMENTO": "2000-01-01",
    })
    columns, rows = exporter.montar_exportacao_simples(conn, "PESSOAS")

    caminho = tmp_path / "export.csv"
    exporter.exportar_csv(columns, rows, str(caminho))

    conteudo = caminho.read_text(encoding="utf-8-sig")
    primeira_linha = conteudo.splitlines()[0]
    assert "DATA_DE_NASCIMENTO" in primeira_linha
    assert "DATA DE NASCIMENTO" not in primeira_linha


def test_exportar_mesclado_troca_traco_por_underscore_no_cabecalho(conn, tmp_path):
    id_empresa = records.create_record(conn, "EMPRESAS", {"SIGLA": "ABC", "EMPRESA": "Empresa ABC"})
    records.create_record(conn, "PESSOAS", {"ID_EMPRESA": id_empresa, "CARGO": "Presidente", "NOME": "Ana"})
    columns, rows = exporter.montar_exportacao_mesclada(
        conn, "PESSOAS", valores_agrupador=["Presidente"], campo_agrupador="CARGO",
        campos_por_valor={"Presidente": ["NOME"]},
    )

    caminho = tmp_path / "mesclado.csv"
    exporter.exportar_csv(columns, rows, str(caminho))

    primeira_linha = caminho.read_text(encoding="utf-8-sig").splitlines()[0]
    assert "Presidente_NOME" in primeira_linha
    assert "Presidente - NOME" not in primeira_linha


def test_exportar_csv_desambigua_colunas_que_colidem_apos_normalizar(tmp_path):
    """"Diretor - Tecnico" e "Diretor Tecnico" sao colunas DIFERENTES, mas
    as duas normalizam pro mesmo texto ("Diretor_Tecnico") -- o arquivo nao
    pode ter duas colunas com o mesmo titulo (uma ferramenta de mala direta
    so enxergaria uma das duas, perdendo a outra silenciosamente)."""
    columns = ["Diretor - Tecnico", "Diretor Tecnico"]
    rows = [{"Diretor - Tecnico": "A", "Diretor Tecnico": "B"}]

    caminho = tmp_path / "colisao.csv"
    exporter.exportar_csv(columns, rows, str(caminho))

    primeira_linha = caminho.read_text(encoding="utf-8-sig").splitlines()[0]
    titulos = [t.strip('"') for t in primeira_linha.split(",")]
    assert titulos == ["Diretor_Tecnico", "Diretor_Tecnico_2"]
